"""PPTX document loader using python-pptx."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import ClassVar

from jmeter_ai.core.exceptions import DocumentParseError
from jmeter_ai.core.schemas import DocumentSection, ParsedDocument
from jmeter_ai.document_processing.base import AbstractDocumentLoader


class PptxLoader(AbstractDocumentLoader):
    """Loads .pptx files, extracting slides as sections."""

    supported_extensions: ClassVar[set[str]] = {"pptx"}

    async def load(self, path: Path) -> ParsedDocument:
        """Load a PowerPoint file and extract slide content."""
        await self._validate_file(path)
        return await asyncio.to_thread(self._load_sync, path)

    def _load_sync(self, path: Path) -> ParsedDocument:
        """Synchronous PPTX loading."""
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
        except Exception as exc:
            raise DocumentParseError(f"Failed to parse PPTX: {exc}") from exc

        sections: list[DocumentSection] = []
        raw_parts: list[str] = []

        for i, slide in enumerate(prs.slides, start=1):
            title = ""
            content_parts: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if not text:
                        continue
                    if shape == slide.shapes.title:
                        title = text
                    else:
                        content_parts.append(text)

            slide_content = "\n".join(content_parts)
            heading = title or f"Slide {i}"
            raw_parts.append(f"{heading}\n{slide_content}")

            if slide_content or title:
                sections.append(
                    DocumentSection(
                        heading=heading,
                        content=slide_content,
                        level=1,
                        type="slide",
                    )
                )

        return ParsedDocument(
            source_path=str(path),
            file_type="pptx",
            raw_text="\n\n".join(raw_parts),
            sections=sections,
            metadata={"slide_count": len(prs.slides)},
            page_count=len(prs.slides),
        )
