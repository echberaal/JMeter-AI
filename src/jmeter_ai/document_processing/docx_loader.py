"""DOCX document loader using python-docx."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import ClassVar

from jmeter_ai.core.exceptions import DocumentParseError
from jmeter_ai.core.schemas import DocumentSection, ParsedDocument
from jmeter_ai.document_processing.base import AbstractDocumentLoader


class DocxLoader(AbstractDocumentLoader):
    """Loads .docx files, extracting paragraphs with style info and tables."""

    supported_extensions: ClassVar[set[str]] = {"docx"}

    async def load(self, path: Path) -> ParsedDocument:
        """Load a DOCX file and extract structured content."""
        await self._validate_file(path)
        return await asyncio.to_thread(self._load_sync, path)

    def _load_sync(self, path: Path) -> ParsedDocument:
        """Synchronous DOCX loading."""
        try:
            from docx import Document

            doc = Document(str(path))
        except Exception as exc:
            raise DocumentParseError(f"Failed to parse DOCX: {exc}") from exc

        sections: list[DocumentSection] = []
        raw_parts: list[str] = []
        current_heading = ""
        current_content: list[str] = []
        current_level = 1

        for para in doc.paragraphs:
            style_name = (para.style.name or "").lower() if para.style else ""
            text = para.text.strip()
            if not text:
                continue

            if "heading" in style_name:
                # Flush previous section
                if current_content:
                    sections.append(
                        DocumentSection(
                            heading=current_heading,
                            content="\n".join(current_content),
                            level=current_level,
                            type="body",
                        )
                    )
                    current_content = []
                current_heading = text
                # Extract heading level (e.g., "heading 2" -> 2)
                try:
                    current_level = int(style_name.replace("heading", "").strip())
                except ValueError:
                    current_level = 1
            else:
                current_content.append(text)
            raw_parts.append(text)

        # Flush last section
        if current_content:
            sections.append(
                DocumentSection(
                    heading=current_heading,
                    content="\n".join(current_content),
                    level=current_level,
                    type="body",
                )
            )

        # Extract tables
        table_texts: list[str] = []
        for table in doc.tables:
            rows_text: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows_text.append(" | ".join(cells))
            table_text = "\n".join(rows_text)
            table_texts.append(table_text)
            sections.append(
                DocumentSection(
                    heading="Table",
                    content=table_text,
                    level=1,
                    type="table",
                )
            )

        raw_text = "\n".join(raw_parts)
        if table_texts:
            raw_text += "\n\n" + "\n\n".join(table_texts)

        return ParsedDocument(
            source_path=str(path),
            file_type="docx",
            raw_text=raw_text,
            sections=sections,
            metadata={"table_count": len(doc.tables)},
            page_count=1,
        )
